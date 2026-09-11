import csv
import io
import re
from zipfile import BadZipFile
from datetime import datetime

from sqlalchemy import select

from .ad_models import AdPlan, AdPlanSnapshot, AdTargetRoiSetting

ALIASES = {
    "campaign_id": ["Campaign ID", "campaign_id", "广告计划 ID", "广告计划ID"],
    "plan_name": ["plan_name", "广告计划", "广告计划名称", "计划名称", "campaign_name", "campaign"],
    "spend": ["spend", "消耗", "花费", "cost", "成本"],
    "orders": ["orders", "订单", "成交订单", "purchases", "SKU订单数", "SKU 订单数"],
    "revenue": ["revenue", "销售额", "成交金额", "gmv", "总收入"],
    "actual_roi": ["actual_roi", "实际roi", "实际 ROI", "roi"],
    "budget": ["budget", "预算", "Current budget"], "product_name": ["product_name", "产品", "产品名称"],
    "market_name": ["market_name", "市场", "国家", "市场名称"], "strategy_code": ["strategy_code", "策略", "策略类型"],
    "product_unit_price": ["product_unit_price", "客单价", "产品客单价", "售价"],
}

REMINDER_STATUSES = {"empty_burn"}
EMPTY_BURN_THRESHOLD = 2.0
PLAN_VARIANT_PATTERN = re.compile(r"^([AB])\*")
PLAN_PRODUCT_PATTERN = re.compile(
    r"^[A-Z]\*\d{1,2}\.\d{1,2}(.*?)(?:\d+(?:\.\d+)?(?:\s+\d+(?:\.\d+)?)*)\.*$"
)


def _number(value, default=None):
    if value is None or str(value).strip() == "": return default
    try: return float(str(value).replace(",", "").replace("%", ""))
    except (TypeError, ValueError): return default


def _find(row, key):
    normalized = {re.sub(r"[\s_\-]", "", str(k)).lower(): v for k, v in row.items()}
    for alias in ALIASES[key]:
        value = normalized.get(re.sub(r"[\s_\-]", "", alias).lower())
        if value is not None: return value
    return None


def _has_column(row, key):
    normalized = {re.sub(r"[\s_\-]", "", str(column)).lower() for column in row}
    return any(re.sub(r"[\s_\-]", "", alias).lower() in normalized for alias in ALIASES[key])


def parse_plan_name(plan_name: str) -> tuple[str | None, str | None]:
    """Parse only the A/B marker and product shape observed in the supplied export."""
    variant_match = PLAN_VARIANT_PATTERN.match(plan_name)
    product_match = PLAN_PRODUCT_PATTERN.match(plan_name)
    product_name = product_match.group(1).strip() if product_match else None
    return (variant_match.group(1) if variant_match else None), (product_name or None)


def normalize_product_name(product_name: str) -> str:
    return re.sub(r"[\s\-_.]+", "", product_name).casefold()


def find_target_setting(db, product_name: str | None, strategy_code: str | None):
    if not product_name or strategy_code not in {"A", "B"}:
        return None
    return db.scalar(select(AdTargetRoiSetting).where(
        AdTargetRoiSetting.product_key == normalize_product_name(product_name),
        AdTargetRoiSetting.strategy_code == strategy_code,
    ))


def parse_upload(filename: str, content: bytes):
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix in {"csv", "txt"}:
        text = content.decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(text)))
    if suffix in {"xlsx", "xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("解析 xlsx 需要安装 openpyxl：pip install openpyxl") from exc
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, OSError, ValueError) as exc:
            raise ValueError("Excel 文件无效或已损坏") from exc
        sheet = workbook.active; rows = sheet.iter_rows(values_only=True); headers = [str(x or "") for x in next(rows, ())]
        return [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
    raise ValueError("仅支持 .xlsx、.xlsm 或 .csv 文件")


def validate_import_rows(rows):
    required = {"plan_name": "广告计划名称", "spend": "成本", "orders": "SKU 订单数", "actual_roi": "ROI"}
    missing = [label for key, label in required.items() if not _has_column(rows[0], key)]
    if missing:
        raise ValueError(f"Excel 缺少必需列：{'、'.join(missing)}")
    for index, row in enumerate(rows, start=2):
        if not str(_find(row, "plan_name") or "").strip():
            raise ValueError(f"Excel 第 {index} 行广告计划名称为空")
        values = {key: _number(_find(row, key), None) for key in ("spend", "orders", "actual_roi")}
        empty = [required[key] for key, value in values.items() if value is None]
        if empty:
            raise ValueError(f"Excel 第 {index} 行{'、'.join(empty)}为空或不是数字")
        if values["spend"] < 0 or values["orders"] < 0 or values["actual_roi"] < 0:
            raise ValueError(f"Excel 第 {index} 行成本、SKU 订单数和 ROI 不能为负数")
        if not values["orders"].is_integer():
            raise ValueError(f"Excel 第 {index} 行 SKU 订单数必须是整数")


def evaluate_snapshot(ad_plan, spend, orders, actual_roi):
    if spend is None or orders is None or actual_roi is None or spend < 0 or orders < 0 or actual_roi < 0:
        return "unknown", "历史快照数据不可用", "请重新导入完整的 TikTok Excel"
    if spend == 0:
        return "no_spend", "本次 Excel 报表周期内成本为 0", "建议回 TikTok 后台检查计划是否正常投放"
    if orders == 0:
        if spend > EMPTY_BURN_THRESHOLD:
            return "empty_burn", "成本已超过 $2，但 SKU 订单数为 0", "建议回 TikTok 后台检查对应广告视频，再决定是否关停或重建"
        return "exploring", "计划已有消耗，但成本未超过 $2 且暂未出单", "计划仍在消耗探索中，继续观察"
    if ad_plan.target_roi is None:
        return "target_roi_pending", "该广告计划尚未设置目标 ROI", "请为这条计划设置目标 ROI"
    if actual_roi >= ad_plan.target_roi:
        return "roi_reached", "当前 ROI 已达到目标 ROI", "可以继续关注，根据实际运营策略决定是否提高 ROI"
    return "low_roi", "当前 ROI 低于目标 ROI", "建议回 TikTok 后台查看对应视频数据，进一步判断问题原因"


def latest_snapshot(ad_plan: AdPlan) -> AdPlanSnapshot | None:
    return max(ad_plan.snapshots, key=lambda item: (item.snapshot_at or datetime.min, item.id), default=None)


def is_alert_pending(ad_plan: AdPlan) -> bool:
    """An empty-burn alert is handled once an operation is recorded after its latest snapshot."""
    if ad_plan.current_status != "empty_burn":
        return False
    latest_operation = max(ad_plan.operations, key=lambda item: (item.operated_at or datetime.min, item.id), default=None)
    if latest_operation is None:
        return True
    snapshot = latest_snapshot(ad_plan)
    return snapshot is not None and (latest_operation.operated_at or datetime.min) < (snapshot.snapshot_at or datetime.min)


def recompute_plan_status(ad_plan: AdPlan) -> AdPlanSnapshot | None:
    snapshot = latest_snapshot(ad_plan)
    if snapshot is None:
        return None
    status, reason, recommendation = evaluate_snapshot(
        ad_plan, snapshot.spend, snapshot.orders, snapshot.actual_roi
    )
    snapshot.status = status
    snapshot.reason = reason
    snapshot.recommendation = recommendation
    ad_plan.current_status = status
    return snapshot


def import_rows(db, filename, rows, markets, products):
    from .ad_models import AdImportBatch
    validate_import_rows(rows)
    batch = AdImportBatch(file_name=filename, row_count=len(rows), status="completed"); db.add(batch); db.flush()
    for row in rows:
        name = str(_find(row, "plan_name") or "").strip()
        if not name: continue
        parsed_strategy, parsed_product = parse_plan_name(name)
        product_name = str(_find(row, "product_name") or parsed_product or "").strip()
        market_name = str(_find(row, "market_name") or "").strip()
        campaign_id = str(_find(row, "campaign_id") or "").strip() or None
        product = next((p for p in products if p.name == product_name), None)
        market = next((m for m in markets if m.name == market_name or m.country_code == market_name), None)
        plan = db.scalar(select(AdPlan).where(AdPlan.platform_campaign_id == campaign_id)) if campaign_id else None
        if plan is None:
            plan = db.scalar(select(AdPlan).where(AdPlan.plan_name == name))
        supplied_strategy = str(_find(row, "strategy_code") or "").strip().upper()
        strategy_code = supplied_strategy if supplied_strategy in {"A", "B"} else parsed_strategy
        imported_price = _number(_find(row, "product_unit_price"), None)
        if not plan:
            plan = AdPlan(platform_campaign_id=campaign_id, plan_name=name, imported_product_name=product_name or None,
                          product_id=product.id if product else None, market_id=market.id if market else None,
                          strategy_code=strategy_code, product_unit_price=imported_price)
            db.add(plan); db.flush()
        else:
            if campaign_id and plan.platform_campaign_id is None:
                plan.platform_campaign_id = campaign_id
            if product_name:
                plan.imported_product_name = product_name
            if plan.product_id is None and product:
                plan.product_id = product.id
            if plan.market_id is None and market:
                plan.market_id = market.id
            if strategy_code:
                plan.strategy_code = strategy_code
        if plan.product_unit_price is None and imported_price is not None:
            plan.product_unit_price = imported_price
        target_setting = find_target_setting(db, product_name or plan.imported_product_name, strategy_code or plan.strategy_code)
        if target_setting:
            plan.target_roi = target_setting.target_roi
        spend = _number(_find(row, "spend"), None)
        orders_value = _number(_find(row, "orders"), None)
        orders = int(orders_value) if orders_value is not None else None
        revenue = _number(_find(row, "revenue"), None)
        roi = _number(_find(row, "actual_roi"), None)
        status, reason, recommendation = evaluate_snapshot(plan, spend, orders, roi)
        snapshot = AdPlanSnapshot(ad_plan_id=plan.id, import_batch_id=batch.id, spend=spend, orders=orders, revenue=revenue, actual_roi=roi, budget=_number(_find(row, "budget"), None), status=status, reason=reason, recommendation=recommendation)
        db.add(snapshot); plan.current_status = status
    return batch
