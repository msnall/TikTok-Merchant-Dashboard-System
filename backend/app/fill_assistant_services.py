import io
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .ad_services import _find, _has_column, _number, normalize_product_name, parse_plan_name, parse_upload


CAMPAIGN_REQUIRED_COLUMNS = {
    "plan_name": "广告计划名称",
    "spend": "成本",
    "orders": "SKU 订单数",
    "revenue": "总收入",
    "actual_roi": "ROI",
}

AFFILIATE_STANDARD_COMMISSION = "预计标准佣金付款"
AFFILIATE_STORE_AD_COMMISSION = "预计店铺广告佣金付款"
AFFILIATE_CURRENCY = "货币单位"
SALES_SKU_COLUMN = "Seller SKU"
SALES_AMOUNT_COLUMN = "SKU Subtotal After Discount"
PRODUCT_CODE_SKU_COLUMN = "sellersku"
PRODUCT_CODE_NAME_COLUMN = "名称"


def _strict_number(row, key, label, row_number, *, integer=False):
    raw_value = _find(row, key)
    value = _number(raw_value, None)
    if value is None:
        raise ValueError(f"Excel 第 {row_number} 行{label}为空或不是数字")
    if value < 0:
        raise ValueError(f"Excel 第 {row_number} 行{label}不能为负数")
    if integer and not value.is_integer():
        raise ValueError(f"Excel 第 {row_number} 行{label}必须是整数")
    return int(value) if integer else value


def analyze_product_campaign(filename, content, target_settings, known_products=()):
    rows = parse_upload(filename, content)
    if not rows:
        raise ValueError("文件没有可计算的数据")
    missing = [label for key, label in CAMPAIGN_REQUIRED_COLUMNS.items() if not _has_column(rows[0], key)]
    if missing:
        raise ValueError(f"Excel 缺少必需列：{'、'.join(missing)}")

    configured_products = [setting.product_name for setting in target_settings]
    product_candidates = list(known_products) + configured_products
    target_map = {
        (setting.product_key, setting.strategy_code): float(setting.target_roi)
        for setting in target_settings
    }
    grouped = defaultdict(list)
    zero_cost_rows = 0

    for row_number, row in enumerate(rows, start=2):
        campaign_name = str(_find(row, "plan_name") or "").strip()
        if not campaign_name:
            raise ValueError(f"Excel 第 {row_number} 行广告计划名称为空")
        strategy_code, product_name = parse_plan_name(campaign_name, product_candidates)
        if not strategy_code or not product_name:
            raise ValueError(f"Excel 第 {row_number} 行广告计划名称无法识别：{campaign_name}")

        cost = _strict_number(row, "spend", "成本", row_number)
        orders = _strict_number(row, "orders", "SKU 订单数", row_number, integer=True)
        revenue = _strict_number(row, "revenue", "总收入", row_number)
        actual_roi = _strict_number(row, "actual_roi", "ROI", row_number)
        budget_raw = _find(row, "budget")
        current_budget = None
        if budget_raw is not None and str(budget_raw).strip() != "":
            current_budget = _strict_number(row, "budget", "Current budget", row_number)

        if cost == 0:
            zero_cost_rows += 1
            continue

        target_roi = target_map.get((normalize_product_name(product_name), strategy_code))
        payout_eligible = orders > 21
        payout_roi = target_roi * 0.9 if target_roi is not None else None
        if not payout_eligible:
            payout_amount = 0.0
            payout_status = "not_eligible"
        elif target_roi is None:
            payout_amount = None
            payout_status = "target_roi_pending"
        else:
            payout_amount = max(cost - revenue / payout_roi, 0.0)
            payout_status = "calculated"

        grouped[product_name].append({
            "row_number": row_number,
            "campaign_name": campaign_name,
            "strategy_code": strategy_code,
            "product_name": product_name,
            "cost": cost,
            "orders": orders,
            "revenue": revenue,
            "actual_roi": actual_roi,
            "current_budget": current_budget,
            "target_roi": target_roi,
            "payout_roi": payout_roi,
            "payout_amount": payout_amount,
            "payout_eligible": payout_eligible,
            "payout_status": payout_status,
        })

    products = []
    for product_name, details in grouped.items():
        missing_count = sum(item["payout_status"] == "target_roi_pending" for item in details)
        products.append({
            "product_name": product_name,
            "plan_count": len(details),
            "total_revenue": sum(item["revenue"] for item in details),
            "total_consumption": sum(item["cost"] for item in details),
            "estimated_payout": None if missing_count else sum(item["payout_amount"] or 0 for item in details),
            "missing_target_roi_count": missing_count,
            "details": details,
        })
    products.sort(key=lambda item: (-item["total_consumption"], item["product_name"]))
    details = [item for product in products for item in product["details"]]
    return {
        "filename": filename,
        "total_rows": len(rows),
        "zero_cost_rows": zero_cost_rows,
        "effective_rows": len(details),
        "product_count": len(products),
        "total_consumption": sum(item["cost"] for item in details),
        "payout_eligible_plans": sum(item["payout_eligible"] for item in details),
        "calculated_payout_plans": sum(item["payout_status"] == "calculated" for item in details),
        "missing_target_roi_plans": sum(item["payout_status"] == "target_roi_pending" for item in details),
        "products": products,
    }


def _normalized_headers(row):
    return {str(key or "").strip(): key for key in row}


def _affiliate_amount(row, source_key, label, row_number):
    raw_value = row.get(source_key)
    if raw_value is None or str(raw_value).strip() == "":
        return Decimal("0"), True
    normalized = str(raw_value).strip().replace(",", "")
    try:
        return Decimal(normalized), False
    except InvalidOperation as exc:
        raise ValueError(f"联盟费用表第 {row_number} 行“{label}”不是有效数字：{raw_value}") from exc


def analyze_affiliate_fees(filename, content):
    rows = parse_upload(filename, content)
    if not rows:
        raise ValueError("联盟费用表文件没有有效数据")
    headers = _normalized_headers(rows[0])
    missing = [
        label for label in (AFFILIATE_STANDARD_COMMISSION, AFFILIATE_STORE_AD_COMMISSION)
        if label not in headers
    ]
    if missing:
        raise ValueError(f"联盟费用表缺少必要字段：{'、'.join(missing)}")

    standard_total = Decimal("0")
    store_ad_total = Decimal("0")
    standard_empty_rows = 0
    store_ad_empty_rows = 0
    currencies = set()
    for row_number, row in enumerate(rows, start=2):
        standard, standard_empty = _affiliate_amount(
            row, headers[AFFILIATE_STANDARD_COMMISSION], AFFILIATE_STANDARD_COMMISSION, row_number
        )
        store_ad, store_ad_empty = _affiliate_amount(
            row, headers[AFFILIATE_STORE_AD_COMMISSION], AFFILIATE_STORE_AD_COMMISSION, row_number
        )
        standard_total += standard
        store_ad_total += store_ad
        standard_empty_rows += standard_empty
        store_ad_empty_rows += store_ad_empty
        currency_key = headers.get(AFFILIATE_CURRENCY)
        currency = str(row.get(currency_key) or "").strip() if currency_key else ""
        if currency:
            currencies.add(currency)
    if len(currencies) > 1:
        raise ValueError(f"联盟费用表包含多个货币单位：{'、'.join(sorted(currencies))}")
    return {
        "filename": filename,
        "total_rows": len(rows),
        "standard_commission_total": float(standard_total),
        "store_ad_commission_total": float(store_ad_total),
        "total_affiliate_fee": float(standard_total + store_ad_total),
        "currency": next(iter(currencies), None),
        "standard_empty_rows": standard_empty_rows,
        "store_ad_empty_rows": store_ad_empty_rows,
    }


def _column_key(value):
    return "".join(str(value or "").split()).replace("_", "").replace("-", "").casefold()


def _required_source_columns(row, required, source_label):
    headers = {_column_key(key): key for key in row}
    missing = [label for label in required if _column_key(label) not in headers]
    if missing:
        raise ValueError(f"{source_label}缺少必要字段：{'、'.join(missing)}")
    return {label: headers[_column_key(label)] for label in required}


def parse_product_code_mappings(filename, content):
    rows = parse_upload(filename, content)
    columns = None
    if rows:
        try:
            columns = _required_source_columns(
                rows[0], (PRODUCT_CODE_SKU_COLUMN, PRODUCT_CODE_NAME_COLUMN), "产品编码文件"
            )
        except ValueError:
            columns = None

    # Some product-code workbooks use one product per worksheet. In that
    # format the worksheet title is ``Seller SKU + 产品名称`` and the first
    # sheet is a non-data rule sheet. Keep accepting the regular two-column
    # format above, then fall back to worksheet-title parsing for xlsx/xlsm.
    if columns is None and filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, OSError, ValueError) as exc:
            raise ValueError("产品编码 Excel 文件无效或已损坏") from exc
        rows = []
        for sheet in workbook.worksheets:
            title = str(sheet.title or "").strip()
            match = re.match(r"^(.+?)([\u4e00-\u9fff].*)$", title)
            if not match or not re.search(r"\d", match.group(1)):
                continue
            rows.append({PRODUCT_CODE_SKU_COLUMN: match.group(1).strip(), PRODUCT_CODE_NAME_COLUMN: match.group(2).strip()})
        if not rows:
            raise ValueError("产品编码文件缺少必要字段：sellersku、名称")
        columns = {PRODUCT_CODE_SKU_COLUMN: PRODUCT_CODE_SKU_COLUMN, PRODUCT_CODE_NAME_COLUMN: PRODUCT_CODE_NAME_COLUMN}
    if not rows:
        raise ValueError("产品编码文件没有有效数据")
    mappings = {}
    source_rows = {}
    duplicate_rows = 0
    for row_number, row in enumerate(rows, start=2):
        seller_sku = str(row.get(columns[PRODUCT_CODE_SKU_COLUMN]) or "").strip()
        name = str(row.get(columns[PRODUCT_CODE_NAME_COLUMN]) or "").strip()
        if not seller_sku and not name:
            continue
        if not seller_sku:
            raise ValueError(f"产品编码第 {row_number} 行 sellersku 为空")
        if not name:
            raise ValueError(f"产品编码第 {row_number} 行名称为空")
        if seller_sku in mappings:
            if mappings[seller_sku] != name:
                first_row = source_rows[seller_sku]
                raise ValueError(
                    f"产品编码映射冲突：{seller_sku} 在第 {first_row} 行为“{mappings[seller_sku]}”，"
                    f"第 {row_number} 行为“{name}”"
                )
            duplicate_rows += 1
            continue
        mappings[seller_sku] = name
        source_rows[seller_sku] = row_number
    if not mappings:
        raise ValueError("产品编码文件中没有可导入的 sellersku 映射")
    return {
        "source_rows": len(rows),
        "duplicate_rows": duplicate_rows,
        "mappings": mappings,
    }


def _sales_amount(row, source_key, row_number):
    raw_value = row.get(source_key)
    normalized = str(raw_value or "").strip().replace(",", "")
    if not normalized:
        raise ValueError(f"全部订单第 {row_number} 行“{SALES_AMOUNT_COLUMN}”为空")
    try:
        amount = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(
            f"全部订单第 {row_number} 行“{SALES_AMOUNT_COLUMN}”不是有效数字：{raw_value}"
        ) from exc
    if not amount.is_finite():
        raise ValueError(
            f"全部订单第 {row_number} 行“{SALES_AMOUNT_COLUMN}”不是有效数字：{raw_value}"
        )
    return amount


def analyze_sales_data(filename, content, product_code_mappings):
    rows = parse_upload(filename, content)
    if not rows:
        raise ValueError("全部订单文件没有有效数据")
    columns = _required_source_columns(
        rows[0], (SALES_SKU_COLUMN, SALES_AMOUNT_COLUMN), "全部订单文件"
    )
    grouped = defaultdict(Decimal)
    total_amount = Decimal("0")
    blank_sku_rows = 0
    valid_rows = 0
    for row_number, row in enumerate(rows, start=2):
        seller_sku = str(row.get(columns[SALES_SKU_COLUMN]) or "").strip()
        amount = _sales_amount(row, columns[SALES_AMOUNT_COLUMN], row_number)
        if not seller_sku:
            blank_sku_rows += 1
            continue
        grouped[seller_sku] += amount
        total_amount += amount
        valid_rows += 1

    # Aggregate by product name only after SKU totals have been calculated.
    # Unknown SKUs remain separate so one "未匹配" row cannot hide codes.
    merged = {}
    for seller_sku, subtotal in grouped.items():
        name = product_code_mappings.get(seller_sku)
        key = ("name", name) if name is not None else ("sku", seller_sku)
        entry = merged.setdefault(key, {
            "name": name or "未匹配",
            "seller_skus": [],
            "subtotal": Decimal("0"),
            "matched": name is not None,
        })
        entry["seller_skus"].append(seller_sku)
        entry["subtotal"] += subtotal
    result_rows = []
    for entry in sorted(merged.values(), key=lambda item: (-item["subtotal"], item["name"], item["seller_skus"])):
        seller_skus = sorted(entry["seller_skus"])
        result_rows.append({
            "seller_sku": ", ".join(seller_skus),
            "seller_skus": seller_skus,
            "name": entry["name"],
            "subtotal_after_discount": float(entry["subtotal"]),
            "matched": entry["matched"],
        })
    matched_sku_count = sum(1 for seller_sku in grouped if seller_sku in product_code_mappings)
    return {
        "filename": filename,
        "total_rows": len(rows),
        "valid_rows": valid_rows,
        "blank_sku_rows": blank_sku_rows,
        "sku_count": len(grouped),
        "product_count": len(result_rows),
        "matched_sku_count": matched_sku_count,
        "unmatched_sku_count": len(grouped) - matched_sku_count,
        "total_amount": float(total_amount),
        "rows": result_rows,
    }


def parse_payout_config(content):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError) as exc:
        raise ValueError("赔付配置 Excel 文件无效或已损坏") from exc
    sheet = workbook.active
    if sheet.calculate_dimension() == "A1:A1":
        sheet.reset_dimensions()
    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    product_column = None
    target_column = None
    for index, row in enumerate(rows):
        normalized = [str(value or "").replace(" ", "").upper() for value in row]
        if "产品名称" in normalized and "目标ROI" in normalized:
            header_index = index
            product_column = normalized.index("产品名称")
            target_column = normalized.index("目标ROI")
            break
    if header_index is None:
        raise ValueError("赔付配置 Excel 缺少“产品名称”或“目标ROI”表头")

    settings = {}
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        campaign_name = row[product_column] if product_column < len(row) else None
        target_raw = row[target_column] if target_column < len(row) else None
        if campaign_name in (None, "") or target_raw in (None, ""):
            continue
        target_roi = _number(target_raw, None)
        if target_roi is None or target_roi <= 0:
            raise ValueError(f"赔付配置第 {row_number} 行目标 ROI 必须是大于 0 的数字")
        strategy_code, product_name = parse_plan_name(str(campaign_name))
        if not strategy_code or not product_name:
            raise ValueError(f"赔付配置第 {row_number} 行计划名称无法识别：{campaign_name}")
        key = (normalize_product_name(product_name), strategy_code)
        if key in settings and settings[key]["target_roi"] != target_roi:
            raise ValueError(f"赔付配置中 {product_name} · {strategy_code} 存在冲突的目标 ROI")
        settings[key] = {
            "product_name": product_name,
            "product_key": key[0],
            "strategy_code": strategy_code,
            "target_roi": target_roi,
            "source_row": row_number,
        }
    if not settings:
        raise ValueError("赔付配置 Excel 中没有可导入的目标 ROI")
    return list(settings.values())


HEADER_FILL = PatternFill("solid", fgColor="173F3A")
SUBHEADER_FILL = PatternFill("solid", fgColor="DDECE8")


def _write_table(sheet, headers, rows, number_columns=()):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in number_columns:
        for cell in sheet[column][1:]:
            cell.number_format = "0.00"
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(sheet.cell(row=row, column=index).value or "") for row in range(2, sheet.max_row + 1)]
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max(map(len, values)) + 2, 12), 42)


def export_consumption_workbook(result):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "全域消耗"
    rows = [[item.product_name, item.plan_count, item.total_consumption] for item in result.products]
    _write_table(sheet, ["产品名称", "计划数量", "全域总消耗"], rows, ("C",))
    stream = io.BytesIO(); workbook.save(stream); stream.seek(0)
    return stream


def export_payout_workbook(result):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "产品汇总"
    summary_rows = [[
        item.product_name, item.total_revenue, item.total_consumption,
        item.estimated_payout, item.missing_target_roi_count,
    ] for item in result.products]
    _write_table(summary, ["产品名称", "总收入", "全域总消耗", "预估赔付", "待设置目标ROI计划数"], summary_rows, ("B", "C", "D"))

    details = workbook.create_sheet("计划明细")
    detail_rows = []
    for product in result.products:
        for item in product.details:
            detail_rows.append([
                item.strategy_code, item.campaign_name, item.product_name, item.orders,
                item.revenue, item.target_roi, item.payout_roi, item.cost,
                item.payout_amount, item.actual_roi, item.current_budget,
            ])
    _write_table(details, [
        "编号", "广告计划名称", "产品", "SKU订单数", "总收入", "目标ROI",
        "赔付ROI", "实际消耗", "赔付金额", "实际ROI", "Current budget",
    ], detail_rows, ("E", "F", "G", "H", "I", "J", "K"))
    stream = io.BytesIO(); workbook.save(stream); stream.seek(0)
    return stream


def export_affiliate_fee_workbook(result):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "精选联盟费用"
    rows = [
        ["预计标准佣金付款", result.standard_commission_total, result.currency],
        ["预计店铺广告佣金付款", result.store_ad_commission_total, result.currency],
        ["精选联盟总费用", result.total_affiliate_fee, result.currency],
    ]
    _write_table(sheet, ["项目", "金额", "币种"], rows, ("B",))
    sheet["A4"].font = Font(bold=True)
    sheet["B4"].font = Font(bold=True)
    sheet["C4"].font = Font(bold=True)
    sheet["A4"].fill = SUBHEADER_FILL
    sheet["B4"].fill = SUBHEADER_FILL
    sheet["C4"].fill = SUBHEADER_FILL
    stream = io.BytesIO(); workbook.save(stream); stream.seek(0)
    return stream


def export_sales_data_workbook(result):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售数据"
    rows = [
        [item.seller_sku, item.name, item.subtotal_after_discount]
        for item in sorted(
            result.rows,
            key=lambda item: (-item.subtotal_after_discount, item.seller_sku),
        )
    ]
    _write_table(
        sheet,
        ["Seller SKU", "名称", "SKU Subtotal After Discount"],
        rows,
        ("C",),
    )
    stream = io.BytesIO(); workbook.save(stream); stream.seek(0)
    return stream
