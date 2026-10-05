"""Read the demo template without treating its inputs as verified platform exports."""

from pathlib import Path

from openpyxl import load_workbook

from .rule_engine import get_roi_policy, parse_campaign_name_roi


CAMPAIGN_FIELDS = {
    "Campaign ID": "campaign_id", "广告计划名称": "campaign_name",
    "产品名": "product", "目标ROI": "target_roi", "成本": "spend",
    "订单": "orders", "总收入": "revenue", "当前ROI": "current_roi",
}
VIDEO_FIELDS = {
    "Video ID": "video_id", "Campaign ID": "campaign_id", "报告日期": "report_date",
    "曝光": "impressions", "点击": "clicks", "CTR": "ctr", "转化订单": "orders",
    "CVR": "cvr", "官方完播率": "official_completion_rate",
    "视频消耗": "spend", "视频收入": "revenue",
}


def read_sheet_rows(path: str | Path, sheet_name: str, header_row: int, data_row: int) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = [str(value or "").strip() for value in next(sheet.iter_rows(
            min_row=header_row, max_row=header_row, values_only=True))]
        rows = []
        for values in sheet.iter_rows(min_row=data_row, values_only=True):
            if values and any(value is not None for value in values):
                rows.append(dict(zip(headers, values)))
        return rows
    finally:
        workbook.close()


def normalize_campaign_row(raw: dict, source_type: str) -> dict:
    if source_type not in {"uploaded_source", "manual_entry", "synthetic_demo"}:
        raise ValueError("无效 source_type")
    if raw.get("source_type") and raw["source_type"] != source_type:
        raise ValueError("原始行与指定 source_type 不一致")
    normalized = {target: raw.get(source) for source, target in CAMPAIGN_FIELDS.items()}
    parsed = parse_campaign_name_roi(str(normalized.get("campaign_name") or ""))
    normalized["link_type"] = parsed["link_type"]
    normalized["roi_policy"] = get_roi_policy(parsed["link_type"])
    normalized["target_roi"] = parsed["target_roi"] if parsed["target_roi"] is not None else normalized["target_roi"]
    normalized["source_type"] = source_type
    normalized["currency"] = raw.get("币种")
    return normalized


def normalize_video_row(raw: dict, source_type: str, source_document: str,
                        official_field_confirmed: bool = False) -> dict:
    if source_type not in {"uploaded_source", "manual_entry", "synthetic_demo"}:
        raise ValueError("无效 source_type")
    if raw.get("source_type") and raw["source_type"] != source_type:
        raise ValueError("原始行与指定 source_type 不一致")
    if source_type != "synthetic_demo" and official_field_confirmed and not source_document.strip():
        raise ValueError("确认真实官方字段时必须提供来源文档")
    normalized = {target: raw.get(source) for source, target in VIDEO_FIELDS.items()}
    if source_type != "synthetic_demo" and not official_field_confirmed:
        normalized["official_completion_rate"] = None
    normalized.update({
        "source_type": source_type,
        "source_document": source_document,
        "source_field_name": "官方完播率" if normalized["official_completion_rate"] is not None else None,
    })
    return normalized
