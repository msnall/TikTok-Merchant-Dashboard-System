import io
import csv

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import Base, engine
from app.main import app


client = TestClient(app)


@pytest.fixture(autouse=True)
def clean_database():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    yield


def workbook_bytes(headers, rows, *, leading_row=None):
    workbook = Workbook()
    sheet = workbook.active
    if leading_row:
        sheet.append(leading_row)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def campaign_workbook(rows):
    return workbook_bytes(
        ["Campaign ID", "广告计划名称", "成本", "Current budget", "SKU 订单数", "总收入", "ROI"],
        rows,
    )


def affiliate_csv(rows, headers=None):
    stream = io.StringIO()
    fieldnames = headers or ["订单 ID", "货币单位", "预计标准佣金付款", "预计店铺广告佣金付款"]
    writer = csv.DictWriter(stream, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def create_roi_settings(include_d=True):
    values = {"A": 3.2, "B": 4.2, "C": 3.0, "D": 2.7}
    if not include_d:
        values.pop("D")
    for code, target in values.items():
        response = client.post("/api/ads/target-roi-settings", json={
            "product_name": "微波炉蒸蛋器", "strategy_code": code, "target_roi": target,
        })
        assert response.status_code == 201


def sample_campaign():
    return campaign_workbook([
        ["1", "A*9.9微波炉蒸蛋器0.82", "10", "100", 21, "10", "1.0"],
        ["2", "B*9.10微波炉蒸蛋器0.82", "20", "100", 22, "18", "0.9"],
        ["3", "C*9.10微波炉蒸蛋器1.23", "10", "100", 30, "100", "10"],
        ["4", "D*9.10微波炉蒸蛋器1.23", "0", "100", 40, "0", "0"],
        ["5", "D*9.11微波炉蒸蛋器1.23.", "5", "100", 25, "0", "0"],
    ])


def test_full_consumption_and_payout_boundaries():
    create_roi_settings()
    response = client.post("/api/fill-assistant/analyze", files={
        "file": ("campaign.xlsx", sample_campaign(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    })
    assert response.status_code == 200
    result = response.json()
    assert result["total_rows"] == 5
    assert result["zero_cost_rows"] == 1
    assert result["effective_rows"] == 4
    assert result["product_count"] == 1
    assert result["total_consumption"] == 45
    assert result["payout_eligible_plans"] == 3
    product = result["products"][0]
    assert product["product_name"] == "微波炉蒸蛋器"
    assert product["plan_count"] == 4
    details = {item["strategy_code"]: item for item in product["details"]}
    assert details["A"]["payout_status"] == "not_eligible"
    assert details["A"]["payout_amount"] == 0
    assert details["B"]["payout_roi"] == pytest.approx(3.78)
    assert details["B"]["payout_amount"] == pytest.approx(20 - 18 / 3.78)
    assert details["C"]["payout_amount"] == 0
    assert details["D"]["payout_amount"] == 5
    assert product["estimated_payout"] == pytest.approx((20 - 18 / 3.78) + 5)


def test_missing_roi_marks_only_eligible_row_pending():
    create_roi_settings(include_d=False)
    result = client.post("/api/fill-assistant/analyze", files={"file": ("campaign.xlsx", sample_campaign())}).json()
    product = result["products"][0]
    assert product["estimated_payout"] is None
    assert product["missing_target_roi_count"] == 1
    assert result["missing_target_roi_plans"] == 1
    assert next(item for item in product["details"] if item["strategy_code"] == "D")["payout_status"] == "target_roi_pending"


def test_payout_config_import_upserts_a_b_c_d():
    content = workbook_bytes(
        [None, "产品名称", "总收入", "目标ROI", "赔付ROI", "实际消耗", "赔付金额"],
        [
            [None, "D*9.10微波炉蒸蛋器1.23", 0, 2.7, None, 0, None],
            [None, "A*9.10微波炉蒸蛋器0.82", 0, 3.2, None, 0, None],
            [None, "C*9.10微波炉蒸蛋器1.23", 0, 3.0, None, 0, None],
            [None, "B*9.10微波炉蒸蛋器0.82", 0, 4.2, None, 0, None],
            [None, "泰国", None, None, None, None, None],
        ],
        leading_row=[None, "填写", "填写", "填写", "/", "填写", "/"],
    )
    imported = client.post("/api/fill-assistant/import-payout-config", files={"file": ("赔付.xlsx", content)})
    assert imported.status_code == 200
    assert imported.json()["created"] == 4
    settings = client.get("/api/ads/target-roi-settings").json()
    assert {(item["strategy_code"], item["target_roi"]) for item in settings} == {
        ("A", 3.2), ("B", 4.2), ("C", 3.0), ("D", 2.7),
    }
    imported_again = client.post("/api/fill-assistant/import-payout-config", files={"file": ("赔付.xlsx", content)})
    assert imported_again.json()["created"] == 0
    assert imported_again.json()["updated"] == 4


def test_exports_match_analyzed_values():
    create_roi_settings()
    analysis = client.post("/api/fill-assistant/analyze", files={"file": ("campaign.xlsx", sample_campaign())}).json()
    consumption = client.post("/api/fill-assistant/export-consumption", json=analysis)
    assert consumption.status_code == 200
    workbook = load_workbook(io.BytesIO(consumption.content), data_only=True)
    assert workbook.sheetnames == ["全域消耗"]
    assert workbook["全域消耗"]["C2"].value == 45
    payout = client.post("/api/fill-assistant/export-payout", json=analysis)
    assert payout.status_code == 200
    workbook = load_workbook(io.BytesIO(payout.content), data_only=True)
    assert workbook.sheetnames == ["产品汇总", "计划明细"]
    assert workbook["产品汇总"]["D2"].value == pytest.approx(analysis["products"][0]["estimated_payout"])


@pytest.mark.parametrize("content, expected", [
    (b"not a workbook", "无效或已损坏"),
    (campaign_workbook([["1", "A*9.10微波炉蒸蛋器0.82", "bad", 10, 1, 2, 2]]), "成本为空或不是数字"),
    (campaign_workbook([["1", "无法解析的计划", 1, 10, 1, 2, 2]]), "广告计划名称无法识别"),
])
def test_analysis_rejects_invalid_input(content, expected):
    response = client.post("/api/fill-assistant/analyze", files={"file": ("campaign.xlsx", content)})
    assert response.status_code == 400
    assert expected in response.json()["detail"]


def test_analysis_rejects_missing_required_column():
    content = workbook_bytes(["广告计划名称", "成本", "SKU 订单数", "ROI"], [["A*9.10产品0.8", 1, 1, 1]])
    response = client.post("/api/fill-assistant/analyze", files={"file": ("campaign.xlsx", content)})
    assert response.status_code == 400
    assert "总收入" in response.json()["detail"]


def test_payout_config_rejects_broken_or_missing_headers():
    broken = client.post("/api/fill-assistant/import-payout-config", files={"file": ("赔付.xlsx", b"broken")})
    assert broken.status_code == 400
    assert "无效或已损坏" in broken.json()["detail"]
    missing = workbook_bytes(["产品名称", "其他字段"], [["A*9.10产品0.8", 2]])
    response = client.post("/api/fill-assistant/import-payout-config", files={"file": ("赔付.xlsx", missing)})
    assert response.status_code == 400
    assert "目标ROI" in response.json()["detail"]


def test_affiliate_fee_analysis_sums_columns_and_treats_blanks_as_zero():
    content = affiliate_csv([
        {"订单 ID": "1", "货币单位": "VND", "预计标准佣金付款": "7500", "预计店铺广告佣金付款": ""},
        {"订单 ID": "2", "货币单位": "VND", "预计标准佣金付款": "", "预计店铺广告佣金付款": "2,500"},
        {"订单 ID": "3", "货币单位": "VND", "预计标准佣金付款": "12.35", "预计店铺广告佣金付款": "1.15"},
    ])
    response = client.post("/api/fill-assistant/affiliate-fees/analyze", files={"file": ("affiliate.csv", content, "text/csv")})
    assert response.status_code == 200
    result = response.json()
    assert result["total_rows"] == 3
    assert result["currency"] == "VND"
    assert result["standard_commission_total"] == pytest.approx(7512.35)
    assert result["store_ad_commission_total"] == pytest.approx(2501.15)
    assert result["total_affiliate_fee"] == pytest.approx(10013.5)
    assert result["standard_empty_rows"] == 1
    assert result["store_ad_empty_rows"] == 1


def test_affiliate_fee_export_matches_analysis():
    content = affiliate_csv([
        {"订单 ID": "1", "货币单位": "VND", "预计标准佣金付款": "100", "预计店铺广告佣金付款": "25"},
    ])
    analysis = client.post("/api/fill-assistant/affiliate-fees/analyze", files={"file": ("affiliate.csv", content)}).json()
    response = client.post("/api/fill-assistant/affiliate-fees/export", json=analysis)
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=True)
    sheet = workbook["精选联盟费用"]
    assert sheet["A2"].value == "预计标准佣金付款"
    assert sheet["B2"].value == 100
    assert sheet["B3"].value == 25
    assert sheet["B4"].value == 125
    assert sheet["C4"].value == "VND"


@pytest.mark.parametrize("filename, content, expected", [
    ("empty.csv", b"", "文件没有有效数据"),
    ("broken.xlsx", b"not a workbook", "无效或已损坏"),
    ("missing.csv", affiliate_csv([{"预计标准佣金付款": "1"}], ["预计标准佣金付款"]), "预计店铺广告佣金付款"),
    ("invalid.csv", affiliate_csv([{
        "订单 ID": "1", "货币单位": "VND", "预计标准佣金付款": "abc", "预计店铺广告佣金付款": "1",
    }]), "第 2 行"),
])
def test_affiliate_fee_rejects_invalid_files(filename, content, expected):
    response = client.post("/api/fill-assistant/affiliate-fees/analyze", files={"file": (filename, content)})
    assert response.status_code == 400
    assert expected in response.json()["detail"]


def sales_mapping_workbook(rows):
    return workbook_bytes(["sellersku", "名称"], rows)


def sales_mapping_sheet_workbook(sheet_titles):
    workbook = Workbook()
    workbook.active.title = "编码规则"
    for title in sheet_titles:
        sheet = workbook.create_sheet(title)
        sheet["A1"] = "款式编码"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def sales_orders_csv(rows, headers=None):
    stream = io.StringIO()
    fieldnames = headers or ["Seller SKU", "SKU Subtotal After Discount", "Order Status"]
    writer = csv.DictWriter(stream, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def test_sales_data_groups_by_sku_then_maps_and_sorts():
    mapping = client.post("/api/fill-assistant/sales/product-codes", files={
        "file": ("codes.xlsx", sales_mapping_workbook([["A", "同名"], ["B", "同名"]]))
    })
    assert mapping.status_code == 200
    analysis = client.post("/api/fill-assistant/sales/analyze", files={
        "file": ("orders.csv", sales_orders_csv([
            {"Seller SKU": "A", "SKU Subtotal After Discount": "2", "Order Status": "Cancelled"},
            {"Seller SKU": "A", "SKU Subtotal After Discount": "3", "Order Status": "Completed"},
            {"Seller SKU": "B", "SKU Subtotal After Discount": "4", "Order Status": "Cancelled"},
            {"Seller SKU": "X", "SKU Subtotal After Discount": "9", "Order Status": "Completed"},
        ]))
    })
    assert analysis.status_code == 200
    result = analysis.json()
    assert [(row["seller_sku"], row["name"], row["subtotal_after_discount"]) for row in result["rows"]] == [
        ("A, B", "同名", 9), ("X", "未匹配", 9)
    ]
    assert result["total_rows"] == 4
    assert result["total_amount"] == 18
    assert result["matched_sku_count"] == 2


def test_sales_data_excludes_blank_sku_and_rejects_invalid_amount():
    response = client.post("/api/fill-assistant/sales/analyze", files={
        "file": ("orders.csv", sales_orders_csv([
            {"Seller SKU": "", "SKU Subtotal After Discount": "2", "Order Status": "Completed"},
            {"Seller SKU": "A", "SKU Subtotal After Discount": "abc", "Order Status": "Completed"},
        ]))
    })
    assert response.status_code == 400
    assert "第 3 行" in response.json()["detail"]

    valid = client.post("/api/fill-assistant/sales/analyze", files={
        "file": ("orders.csv", sales_orders_csv([
            {"Seller SKU": "", "SKU Subtotal After Discount": "2", "Order Status": "Completed"},
            {"Seller SKU": "A", "SKU Subtotal After Discount": "1,234.50", "Order Status": "Completed"},
        ]))
    }).json()
    assert valid["blank_sku_rows"] == 1
    assert valid["valid_rows"] == 1
    assert valid["total_amount"] == 1234.5


def test_sales_mapping_conflict_is_rejected_and_export_matches_result():
    conflict = client.post("/api/fill-assistant/sales/product-codes", files={
        "file": ("codes.xlsx", sales_mapping_workbook([["A", "甲"], ["A", "乙"]]))
    })
    assert conflict.status_code == 400
    assert "映射冲突" in conflict.json()["detail"]

    client.post("/api/fill-assistant/sales/product-codes", files={
        "file": ("codes.xlsx", sales_mapping_workbook([["A", "甲"], ["A", "甲"]]))
    })
    result = client.post("/api/fill-assistant/sales/analyze", files={
        "file": ("orders.csv", sales_orders_csv([
            {"Seller SKU": "A", "SKU Subtotal After Discount": "5", "Order Status": "Completed"},
        ]))
    }).json()
    exported = client.post("/api/fill-assistant/sales/export", json=result)
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content), data_only=True)["销售数据"]
    assert [sheet.cell(1, col).value for col in range(1, 4)] == ["Seller SKU", "名称", "SKU Subtotal After Discount"]
    assert [sheet.cell(2, col).value for col in range(1, 4)] == ["A", "甲", 5]


def test_sales_mapping_accepts_one_product_per_worksheet_format():
    response = client.post("/api/fill-assistant/sales/product-codes", files={
        "file": ("产品编码表.xlsx", sales_mapping_sheet_workbook(["A0001搓脚板", "A0002弹簧双圈钥匙扣"]))
    })
    assert response.status_code == 200
    assert response.json()["unique_mappings"] == 2
    result = client.post("/api/fill-assistant/sales/analyze", files={
        "file": ("orders.csv", sales_orders_csv([
            {"Seller SKU": "A0002", "SKU Subtotal After Discount": "7", "Order Status": "Completed"},
            {"Seller SKU": "A0001", "SKU Subtotal After Discount": "3", "Order Status": "Completed"},
        ]))
    }).json()
    assert [(row["seller_sku"], row["name"]) for row in result["rows"]] == [
        ("A0002", "弹簧双圈钥匙扣"), ("A0001", "搓脚板")
    ]
